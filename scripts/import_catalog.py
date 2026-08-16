import json
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/ubuntu/upload/Catalogo_Normas_Penais_Brasil.xlsx")
DESTINATION = ROOT / "src" / "data" / "catalogoNormas.json"

def rows(ws):
    values = list(ws.iter_rows(values_only=True))
    header = [str(cell).strip() if cell is not None else "" for cell in values[0]]
    return [dict(zip(header, row)) for row in values[1:] if any(cell is not None for cell in row)]

def legal_statements(ws):
    values = list(ws.iter_rows(values_only=True))[1:]
    return [
        {
            "enunciado": str(row[0] or "").strip(),
            "teor": str(row[1] or "").strip(),
            "incidencia": str(row[2] or "").strip(),
        }
        for row in values
        if any(cell is not None for cell in row)
    ]

workbook = load_workbook(SOURCE, data_only=True)
catalogo = rows(workbook["Catalogo"])

payload = {
    "norms": [
        {
            "block": row["Bloco tematico"],
            "norm": row["Norma"],
            "subject": row["Objeto"],
            "nature": row["Natureza"],
            "officialUrl": row["Endereco oficial"],
        }
        for row in catalogo
    ],
    "bindingStatements": legal_statements(workbook["Sumulas vinculantes"]),
    "courtStatements": legal_statements(workbook["Sumulas STF e STJ"]),
    "precedents": legal_statements(workbook["Precedentes vinculantes"]),
    "importedAt": datetime.now(timezone.utc).isoformat(),
}

DESTINATION.parent.mkdir(parents=True, exist_ok=True)
DESTINATION.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Catálogo importado: {len(payload['norms'])} normas, {len(payload['bindingStatements'])} súmulas vinculantes.")
